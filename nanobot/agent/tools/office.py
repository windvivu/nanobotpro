"""Structured Office document tools.

The normal ``read_file`` tool can extract Office text.  These tools add
structured reads and conservative writers for the office-secretary role.
All paths go through the same workspace scope as the filesystem tools.
"""

from __future__ import annotations

import importlib.util
import os
import tempfile
import zipfile
from datetime import date, datetime, time
from itertools import islice
from pathlib import Path
from typing import Any

from nanobot.agent.tools.base import tool_parameters
from nanobot.agent.tools.filesystem import _FsTool
from nanobot.agent.tools.schema import (
    ArraySchema,
    BooleanSchema,
    IntegerSchema,
    ObjectSchema,
    StringSchema,
    tool_parameters_schema,
)

# A spreadsheet or table cell: text, a number, true/false or empty. anyOf rather than a list of
# types: the tool registry casts values by "type" and would turn every number into text.
_CELL: dict[str, Any] = {
    "anyOf": [{"type": "string"}, {"type": "number"}, {"type": "boolean"}, {"type": "null"}],
    "description": "Cell value",
}
_FORMULAS_HELP = ("Write text that starts with '=' as a formula. Only for formulas you wrote yourself: "
                  "text copied from a document, an email or the web must stay text")
# edit_excel loads the whole workbook into memory, many times its size on disk
_EDIT_MAX_BYTES = 20 * 1024 * 1024


def _missing(package: str, extra: str = "office") -> str:
    return f"Error: Office tool requires {package}. Install with: pip install -e \".[{extra}]\""


def _atomic_save(path: Path, save: Any) -> None:
    """Save a generated document beside the destination, then replace it."""
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, raw_tmp = tempfile.mkstemp(prefix=f".{path.stem}-", suffix=path.suffix, dir=path.parent)
    os.close(fd)
    tmp = Path(raw_tmp)
    try:
        save(tmp)
        os.replace(tmp, path)
    finally:
        tmp.unlink(missing_ok=True)


def _cell_text(value: Any) -> str:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if value is None:
        return ""
    return str(value)


def _matrix_schema(description: str) -> ArraySchema:
    return ArraySchema(ArraySchema(_CELL, description="Cell values"), description=description)


def _set_cell(cell: Any, value: Any, formulas: bool) -> None:
    """Text starting with '=' stays text unless formulas were asked for: content taken from a document
    or the web must not run as a formula (e.g. =HYPERLINK sending data out) when the file is opened."""
    cell.value = value
    if not formulas and isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"


def _next_row(ws: Any) -> int:
    """The first row below the data, where ws.append() would write."""
    if ws.max_row == 1 and all(cell.value is None for cell in ws[1]):
        return 1
    return ws.max_row + 1


def _pillow_installed() -> bool:
    return importlib.util.find_spec("PIL") is not None


def _lost_on_save(path: Path) -> list[str]:
    """What openpyxl does not read in this workbook, so that saving it would drop them."""
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()

        def has(prefix: str) -> bool:
            return any(name.startswith(prefix) for name in names)

        def contains(prefix: str, *markers: bytes) -> bool:
            for name in names:
                if not (name.startswith(prefix) and name.endswith(".xml")):
                    continue
                with archive.open(name) as stream:  # in chunks: a sheet can be large once unzipped
                    tail = b""
                    while chunk := stream.read(1 << 20):
                        block = tail + chunk
                        if any(marker in block for marker in markers):
                            return True
                        tail = block[-64:]
            return False

        lost = []
        if contains("xl/drawings/drawing", b"<xdr:sp", b"<xdr:grpSp", b"<xdr:cxnSp"):
            lost.append("shapes or text boxes")
        if has("xl/media/") and not _pillow_installed():
            lost.append("images")
        if has("xl/ctrlProps/") or has("xl/activeX/"):
            lost.append("form controls")
        if has("xl/slicers/") or has("xl/timelines/"):
            lost.append("slicers or timelines")
        if has("xl/embeddings/"):
            lost.append("embedded objects")
        if has("xl/charts/chartEx"):
            lost.append("newer chart types")
        if has("xl/model/"):
            lost.append("a data model")
        if contains("xl/worksheets/sheet", b"sparklineGroup"):
            lost.append("sparklines")
        return lost


@tool_parameters(tool_parameters_schema(
    path=StringSchema("Path to an .xlsx workbook"),
    sheet=StringSchema("Optional sheet name; omit to read all sheets", nullable=True),
    max_rows=IntegerSchema(200, description="Maximum rows per sheet", minimum=1, maximum=2000),
    max_columns=IntegerSchema(50, description="Maximum columns per sheet", minimum=1, maximum=200),
    formulas=BooleanSchema(description="Read formulas instead of calculated values", default=False),
    required=["path"],
))
class ReadExcelTool(_FsTool):
    """Read workbook structure and bounded cell values."""

    @property
    def name(self) -> str:
        return "read_excel"

    @property
    def description(self) -> str:
        return ("Read an .xlsx workbook with sheet names, dimensions and bounded tabular values. A formula "
                "cell shows its last calculated value; files written by these tools have none until they "
                "are opened in Excel, so use formulas=true to see the formulas.")

    @property
    def read_only(self) -> bool:
        return True

    async def execute(self, path: str, sheet: str | None = None, max_rows: int = 200,
                      max_columns: int = 50, formulas: bool = False, **kwargs: Any) -> str:
        try:
            from openpyxl import load_workbook
        except ImportError:
            return _missing("openpyxl")
        try:
            fp = self._resolve(path)
            if fp.suffix.lower() != ".xlsx":
                return "Error: read_excel only supports .xlsx files"
            if not fp.is_file():
                return f"Error: File not found: {path}"
            wb = load_workbook(fp, read_only=True, data_only=not formulas)
            try:
                names = wb.sheetnames
                selected = [sheet] if sheet else names
                unknown = [name for name in selected if name not in names]
                if unknown:
                    return f"Error: Unknown sheet: {unknown[0]} (available: {', '.join(names)})"
                parts: list[str] = [f"Workbook: {fp.name}", f"Sheets: {', '.join(names)}"]
                for name in selected:
                    ws = wb[name]
                    declared = (ws.max_row, ws.max_column)  # the file's own <dimension>: may be missing or wrong
                    # Read the rows themselves: openpyxl stops at <dimension>, which some writers leave out
                    # (TypeError) or get wrong (silently fewer rows)
                    ws.reset_dimensions()
                    rows = list(islice(ws.iter_rows(values_only=True), max_rows + 1))
                    shown = rows[:max_rows]
                    width = max((len(row) for row in shown), default=0)
                    size = ""
                    if None not in declared and declared[0] >= len(shown) and declared[1] >= width:
                        size = f" ({declared[0]} rows x {declared[1]} columns)"
                    parts.append(f"\n--- Sheet: {name}{size} ---")
                    for row in shown:
                        parts.append("\t".join(_cell_text(value) for value in row[:max_columns]).rstrip())
                    if len(rows) > max_rows:
                        parts.append(f"(truncated at {max_rows} rows)")
                    if width > max_columns:
                        parts.append(f"(truncated at {max_columns} columns)")
                return "\n".join(parts)
            finally:
                wb.close()  # read-only mode keeps the file open, locked on Windows, until closed
        except PermissionError as exc:
            return f"Error: {exc}"
        except Exception as exc:
            return f"Error reading Excel workbook: {exc}"


@tool_parameters(tool_parameters_schema(
    path=StringSchema("Destination .xlsx path"),
    headers=ArraySchema(StringSchema("Column header"), description="Header row"),
    rows=_matrix_schema("Rows of cell values"),
    sheet_name=StringSchema("Worksheet name", min_length=1, max_length=31),
    overwrite=BooleanSchema(description="Allow replacing an existing file", default=False),
    formulas=BooleanSchema(description=_FORMULAS_HELP, default=False),
    required=["path", "headers", "rows"],
))
class WriteExcelTool(_FsTool):
    """Create a simple, formatted workbook from tabular data."""

    @property
    def name(self) -> str:
        return "write_excel"

    @property
    def description(self) -> str:
        return "Create an .xlsx workbook with a bold header, freeze panes, filter and sized columns."

    @property
    def exclusive(self) -> bool:
        return True

    async def execute(self, path: str, headers: list[str], rows: list[list[Any]],
                      sheet_name: str = "Sheet1", overwrite: bool = False, formulas: bool = False,
                      **kwargs: Any) -> str:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font
        except ImportError:
            return _missing("openpyxl")
        try:
            fp = self._resolve(path)
            if fp.suffix.lower() != ".xlsx":
                return "Error: write_excel only supports .xlsx files"
            if fp.exists() and not overwrite:
                return f"Error: File already exists: {path}. Set overwrite=true to replace it."
            if not headers:
                return "Error: headers must contain at least one column"
            if len(headers) > 200 or len(rows) > 10000:
                return "Error: workbook exceeds the safe limit (200 columns or 10000 rows)"
            wb = Workbook()
            ws = wb.active
            ws.title = sheet_name[:31]
            for column, header in enumerate(headers, 1):
                _set_cell(ws.cell(row=1, column=column), header, formulas)
            for row_index, row in enumerate(rows, 2):
                for column, value in enumerate(list(row)[:len(headers)], 1):
                    _set_cell(ws.cell(row=row_index, column=column), value, formulas)
            for cell in ws[1]:
                cell.font = Font(bold=True)
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for column in ws.columns:
                letter = column[0].column_letter
                width = min(50, max(10, max((len(_cell_text(c.value)) for c in column), default=0) + 2))
                ws.column_dimensions[letter].width = width
            _atomic_save(fp, wb.save)
            return f"Created Excel workbook: {fp} ({len(rows)} rows, {len(headers)} columns)"
        except PermissionError as exc:
            return f"Error: {exc}"
        except Exception as exc:
            return f"Error writing Excel workbook: {exc}"


@tool_parameters(tool_parameters_schema(
    path=StringSchema("Existing .xlsx path"),
    updates=ArraySchema(ObjectSchema({
        "cell": StringSchema("Cell address, e.g. B4"),
        "value": _CELL,
        "sheet": StringSchema("Optional sheet name", nullable=True),
    }, required=["cell", "value"]), description="Cell updates"),
    append_rows=_matrix_schema("Rows to append"),
    sheet=StringSchema("Default sheet for updates/appends", nullable=True),
    save_as=StringSchema("Write the edited workbook to this new .xlsx path and leave the original unchanged",
                         nullable=True),
    formulas=BooleanSchema(description=_FORMULAS_HELP, default=False),
    required=["path"],
))
class EditExcelTool(_FsTool):
    """Apply bounded cell edits while retaining existing workbook formatting."""

    @property
    def name(self) -> str:
        return "edit_excel"

    @property
    def description(self) -> str:
        return ("Edit selected .xlsx cells or append rows, keeping the formatting, formulas, charts and other "
                "parts openpyxl can read. A workbook with parts it would drop (shapes, form controls, "
                "slicers...) is only edited into a copy: pass save_as.")

    @property
    def exclusive(self) -> bool:
        return True

    async def execute(self, path: str, updates: list[dict[str, Any]] | None = None,
                      append_rows: list[list[Any]] | None = None, sheet: str | None = None,
                      save_as: str | None = None, formulas: bool = False, **kwargs: Any) -> str:
        try:
            from openpyxl import load_workbook
        except ImportError:
            return _missing("openpyxl")
        try:
            fp = self._resolve(path)
            if fp.suffix.lower() != ".xlsx":
                return "Error: edit_excel only supports .xlsx files"
            if not fp.is_file():
                return f"Error: File not found: {path}"
            destination = fp
            if save_as:
                destination = self._resolve(save_as)
                if destination.suffix.lower() != ".xlsx":
                    return "Error: save_as must be an .xlsx path"
                if destination.exists():
                    return f"Error: File already exists: {save_as}. Choose a new name."
            if fp.stat().st_size > _EDIT_MAX_BYTES:
                return (f"Error: {path} is larger than {_EDIT_MAX_BYTES // (1024 * 1024)} MB; "
                        "editing it would take too much memory")
            lost = _lost_on_save(fp)
            if lost and not save_as:
                return (f"Error: {path} contains {', '.join(lost)}, which editing would remove. "
                        "Pass save_as to write the edited copy to a new file and keep the original.")
            wb = load_workbook(fp)
            count = 0
            for update in (updates or []):
                target = update.get("sheet") or sheet or wb.active.title
                if target not in wb.sheetnames:
                    wb.close()
                    return f"Error: Unknown sheet: {target}"
                _set_cell(wb[target][str(update["cell"])], update.get("value"), formulas)
                count += 1
            if append_rows:
                target = sheet or wb.active.title
                if target not in wb.sheetnames:
                    wb.close()
                    return f"Error: Unknown sheet: {target}"
                ws = wb[target]
                row_index = _next_row(ws)
                for row in append_rows:
                    for column, value in enumerate(row, 1):
                        _set_cell(ws.cell(row=row_index, column=column), value, formulas)
                    row_index += 1
                count += len(append_rows)
            if not count:
                wb.close()
                return "Error: Provide updates or append_rows"
            # openpyxl keeps formulas but not their results: Excel recalculates them when it opens the file
            wb.calculation.fullCalcOnLoad = True
            has_formulas = any(cell.data_type == "f" for ws in wb.worksheets for row in ws.iter_rows() for cell in row)
            _atomic_save(destination, wb.save)
            wb.close()
            note = ""
            if lost:
                note += f" Not kept in the copy: {', '.join(lost)}; the original is unchanged."
            if has_formulas:
                note += (" Formula results are recalculated when the file is opened in Excel; until then "
                         "read_excel shows formula cells empty unless formulas=true.")
            return f"Updated Excel workbook: {destination} ({count} change(s)).{note}"
        except PermissionError as exc:
            return f"Error: {exc}"
        except Exception as exc:
            return f"Error editing Excel workbook: {exc}"


@tool_parameters(tool_parameters_schema(
    path=StringSchema("Path to a .docx document"),
    include_tables=BooleanSchema(description="Include table contents", default=True),
    required=["path"],
))
class ReadWordTool(_FsTool):
    """Read paragraphs and tables from a Word document."""

    @property
    def name(self) -> str:
        return "read_word"

    @property
    def description(self) -> str:
        return "Read .docx paragraphs and tables with their order preserved as text."

    @property
    def read_only(self) -> bool:
        return True

    async def execute(self, path: str, include_tables: bool = True, **kwargs: Any) -> str:
        try:
            from docx import Document
            from docx.table import Table
            from docx.text.paragraph import Paragraph
        except ImportError:
            return _missing("python-docx")
        try:
            fp = self._resolve(path)
            if fp.suffix.lower() != ".docx":
                return "Error: read_word only supports .docx files"
            if not fp.is_file():
                return f"Error: File not found: {path}"
            doc = Document(fp)
            parts: list[str] = []
            tables = 0
            for child in doc.element.body.iterchildren():  # paragraphs and tables in document order
                tag = child.tag.rsplit("}", 1)[-1]
                if tag == "p":
                    text = Paragraph(child, doc).text
                    if text.strip():
                        parts.append(text)
                elif tag == "tbl" and include_tables:
                    tables += 1
                    parts.append(f"\n--- Table {tables} ---")
                    parts.extend("\t".join(cell.text for cell in row.cells) for row in Table(child, doc).rows)
            text = "\n".join(parts)
            return text[:200_000] + ("\n(truncated)" if len(text) > 200_000 else "")
        except PermissionError as exc:
            return f"Error: {exc}"
        except Exception as exc:
            return f"Error reading Word document: {exc}"


# One piece of a Word document; write_word adds them in the order given
_BLOCK = ObjectSchema({
    "type": StringSchema("What this block is", enum=["heading", "paragraph", "table"]),
    "text": StringSchema("Text of a heading or a paragraph"),
    "level": IntegerSchema(1, description="Heading level", minimum=1, maximum=9),
    "headers": ArraySchema(StringSchema("Column header"), description="Table header row"),
    "rows": _matrix_schema("Table rows"),
}, required=["type"])


@tool_parameters(tool_parameters_schema(
    path=StringSchema("Destination .docx path"),
    title=StringSchema("Optional document title", nullable=True),
    blocks=ArraySchema(_BLOCK, description="Headings, paragraphs and tables, in document order"),
    overwrite=BooleanSchema(description="Allow replacing an existing file", default=False),
    required=["path"],
))
class WriteWordTool(_FsTool):
    """Create a .docx with headings, paragraphs and tables."""

    @property
    def name(self) -> str:
        return "write_word"

    @property
    def description(self) -> str:
        return "Create a .docx document: an optional title, then headings, paragraphs and tables in the order given."

    @property
    def exclusive(self) -> bool:
        return True

    async def execute(self, path: str, title: str | None = None, blocks: list[dict[str, Any]] | None = None,
                      overwrite: bool = False, **kwargs: Any) -> str:
        try:
            from docx import Document
        except ImportError:
            return _missing("python-docx")
        try:
            fp = self._resolve(path)
            if fp.suffix.lower() != ".docx":
                return "Error: write_word only supports .docx files"
            if fp.exists() and not overwrite:
                return f"Error: File already exists: {path}. Set overwrite=true to replace it."
            doc = Document()
            if title:
                doc.add_heading(title, level=0)
            for block in blocks or []:
                kind = block.get("type")
                if kind == "heading":
                    doc.add_heading(str(block.get("text", "")), level=int(block.get("level") or 1))
                elif kind == "paragraph":
                    doc.add_paragraph(str(block.get("text", "")))
                elif kind == "table":
                    rows = list(block.get("rows") or [])
                    headers = list(block.get("headers") or [])
                    width = len(headers) or max((len(row) for row in rows), default=0)
                    if width == 0:
                        continue
                    table = doc.add_table(rows=1 if headers else 0, cols=width)
                    if headers:
                        for cell, value in zip(table.rows[0].cells, headers):
                            cell.text = _cell_text(value)
                    for row in rows:
                        cells = table.add_row().cells
                        for cell, value in zip(cells, row):
                            cell.text = _cell_text(value)
                else:
                    return f"Error: Unknown block type: {kind!r} (use heading, paragraph or table)"
            _atomic_save(fp, doc.save)
            return f"Created Word document: {fp}"
        except PermissionError as exc:
            return f"Error: {exc}"
        except Exception as exc:
            return f"Error writing Word document: {exc}"


@tool_parameters(tool_parameters_schema(
    path=StringSchema("Path to a .pdf document"),
    pages=StringSchema("Optional page range, for example 1-5", nullable=True),
    max_chars=IntegerSchema(128000, description="Maximum extracted characters", minimum=1000, maximum=500000),
    required=["path"],
))
class ReadPdfTool(_FsTool):
    """Extract text from a PDF with optional page range."""

    @property
    def name(self) -> str:
        return "read_pdf"

    @property
    def description(self) -> str:
        return "Extract text from a PDF; scanned pages require OCR and may return little or no text."

    @property
    def read_only(self) -> bool:
        return True

    async def execute(self, path: str, pages: str | None = None, max_chars: int = 128000, **kwargs: Any) -> str:
        try:
            from pypdf import PdfReader
        except ImportError:
            return _missing("pypdf")
        try:
            fp = self._resolve(path)
            if fp.suffix.lower() != ".pdf":
                return "Error: read_pdf only supports .pdf files"
            if not fp.is_file():
                return f"Error: File not found: {path}"
            reader = PdfReader(str(fp))
            total = len(reader.pages)
            start, end = 0, total
            if pages:
                bits = pages.strip().split("-")
                if len(bits) > 2 or not all(bit.strip().isdigit() for bit in bits):
                    return "Error: Invalid page range; use '1' or '1-5'"
                first, last = int(bits[0]), int(bits[-1])
                if first < 1:
                    return "Error: Pages start at 1"
                start = first - 1
                end = last if len(bits) == 2 else first
                if start >= total or end <= start:
                    return f"Error: Page range is outside the document ({total} pages)"
                end = min(end, total)
            parts = [f"PDF: {fp.name} ({total} pages)"]
            length = len(parts[0])
            for index in range(start, end):
                if length > max_chars:  # enough text already: further pages would only be cut off
                    break
                page = f"--- Page {index + 1} ---\n{reader.pages[index].extract_text() or ''}"
                parts.append(page)
                length += len(page) + 2
            text = "\n\n".join(parts)
            return text[:max_chars] + ("\n(truncated)" if len(text) > max_chars else "")
        except PermissionError as exc:
            return f"Error: {exc}"
        except Exception as exc:
            return f"Error reading PDF: {exc}"


OFFICE_TOOL_CLASSES = (
    ReadExcelTool, WriteExcelTool, EditExcelTool,
    ReadWordTool, WriteWordTool, ReadPdfTool,
)
